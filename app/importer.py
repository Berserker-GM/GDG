"""Import the Google Form export (.xlsx) and an optional USN -> Team mapping file.

Importing is idempotent: members are upserted by USN, members missing from the new export
become inactive, and a changed profile URL wipes that member's stored badges.
"""

from __future__ import annotations

import csv
import logging
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    IMPORT_STATUSES,
    INVALID_URL,
    MISSING_URL,
    UNSUPPORTED_PLATFORM,
    Badge,
    Member,
    Team,
)
from app.urls import check_profile_url

log = logging.getLogger(__name__)

FORM_SHEET = "form responses 1"

FLAG_NOT_REGISTERED = "not_registered"
FLAG_INVALID_EMAIL = "invalid_email"
FLAG_INVALID_COLLEGE_EMAIL = "invalid_college_email"
FLAG_SHARED_PROFILE = "shared_profile"
FLAG_DUPLICATE_SUBMISSION = "duplicate_submission"


class FormImportError(Exception):
    """Raised when the spreadsheet cannot be understood at all."""


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_usn(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


def normalize_year(value: Any) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return str(int(value)) if float(value).is_integer() else str(value)
    text = str(value).strip()
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
    except ValueError:
        pass
    return text


def _parse_timestamp(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = _clean(value)
    if not text:
        return None
    for fmt in ("%m/%d/%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


# Column key -> predicate on the normalised header. Order matters: first match wins per column.
_COLUMN_RULES: list[tuple[str, Any]] = [
    ("timestamp", lambda h: "timestamp" in h),
    ("college_email", lambda h: "college" in h and "email" in h),
    ("email", lambda h: "email" in h),
    ("full_name", lambda h: "name" in h and "team" not in h),
    ("usn", lambda h: h == "usn" or re.search(r"\busn\b", h) is not None),
    ("year", lambda h: "year" in h),
    ("department", lambda h: "department" in h or h == "dept"),
    ("registration", lambda h: "registration" in h or "have you completed" in h),
    ("profile_url", lambda h: "profile" in h and ("link" in h or "url" in h)),
    ("team", lambda h: h == "team" or h.startswith("team ") or h == "team name"),
]
_IGNORED = ("upload", "screenshot", "id card")


def map_headers(headers: list[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, raw in enumerate(headers):
        header = _norm_header(raw)
        if not header or any(word in header for word in _IGNORED) or header.startswith("column "):
            continue
        for key, predicate in _COLUMN_RULES:
            if key not in mapping and predicate(header):
                mapping[key] = index
                break
    missing = [key for key in ("usn", "profile_url") if key not in mapping]
    if missing:
        raise FormImportError(f"Required column(s) not found: {', '.join(missing)}")
    return mapping


@dataclass
class FormRow:
    row_number: int
    timestamp: datetime | None
    usn: str
    full_name: str
    email: str
    college_email: str
    year: str | None
    department: str
    registration: str
    profile_url_raw: str
    team: str


@dataclass
class ImportReport:
    total_rows: int = 0
    skipped_rows: list[dict[str, Any]] = field(default_factory=list)
    unique_students: int = 0
    created: int = 0
    updated: int = 0
    reactivated: int = 0
    deactivated: int = 0
    profile_changed: int = 0
    duplicates_resolved: list[dict[str, Any]] = field(default_factory=list)
    usn_conflicts: list[dict[str, Any]] = field(default_factory=list)
    missing_links: list[dict[str, Any]] = field(default_factory=list)
    invalid_links: list[dict[str, Any]] = field(default_factory=list)
    unsupported_links: list[dict[str, Any]] = field(default_factory=list)
    shared_profiles: list[dict[str, Any]] = field(default_factory=list)
    not_registered: list[dict[str, Any]] = field(default_factory=list)
    invalid_emails: list[dict[str, Any]] = field(default_factory=list)
    teams: int = 0
    team_source: str = "none"
    unassigned: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {k: v for k, v in self.__dict__.items()}


def _read_form_rows(path: str | Path, report: ImportReport) -> tuple[list[FormRow], bool]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next(
            (ws for ws in workbook.worksheets if ws.title.strip().lower() == FORM_SHEET),
            workbook.worksheets[0],
        )
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration as exc:
            raise FormImportError("The sheet is empty") from exc
        cols = map_headers(headers)

        def cell(row: tuple, key: str) -> Any:
            index = cols.get(key)
            return row[index] if index is not None and index < len(row) else None

        result: list[FormRow] = []
        for row_number, row in enumerate(rows, start=2):
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            report.total_rows += 1
            usn = normalize_usn(cell(row, "usn"))
            if not usn:
                report.skipped_rows.append({"row": row_number, "reason": "missing USN"})
                continue
            result.append(
                FormRow(
                    row_number=row_number,
                    timestamp=_parse_timestamp(cell(row, "timestamp")),
                    usn=usn,
                    full_name=_clean(cell(row, "full_name")),
                    email=_clean(cell(row, "email")).replace(" ", ""),
                    college_email=_clean(cell(row, "college_email")).replace(" ", ""),
                    year=normalize_year(cell(row, "year")),
                    department=_clean(cell(row, "department")),
                    registration=_clean(cell(row, "registration")),
                    profile_url_raw=str(cell(row, "profile_url") or "").strip(),
                    team=_clean(cell(row, "team")),
                )
            )
        return result, "team" in cols
    finally:
        workbook.close()


def read_team_mapping(path: str | Path) -> dict[str, str]:
    """Read a USN -> Team mapping from CSV or XLSX (headers matched case-insensitively)."""
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
    else:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            rows = [list(r) for r in csv.reader(fh)]
    if not rows:
        return {}
    headers = [_norm_header(h) for h in rows[0]]
    try:
        usn_col = next(i for i, h in enumerate(headers) if "usn" in h)
        team_col = next(i for i, h in enumerate(headers) if "team" in h)
    except StopIteration as exc:
        raise FormImportError("Teams file needs 'USN' and 'Team' columns") from exc
    mapping: dict[str, str] = {}
    for row in rows[1:]:
        usn = normalize_usn(row[usn_col] if usn_col < len(row) else "")
        team = _clean(row[team_col] if team_col < len(row) else "")
        if usn and team:
            mapping[usn] = team
    return mapping


def _dedupe(rows: list[FormRow], report: ImportReport) -> list[FormRow]:
    by_usn: dict[str, list[FormRow]] = defaultdict(list)
    for row in rows:
        by_usn[row.usn].append(row)
    kept: list[FormRow] = []
    for usn, group in by_usn.items():
        # Latest timestamp wins; rows without a timestamp lose, ties go to the later row.
        group.sort(key=lambda r: (r.timestamp or datetime.min, r.row_number))
        latest = group[-1]
        kept.append(latest)
        if len(group) == 1:
            continue
        profile_ids = {check_profile_url(r.profile_url_raw).profile_id or r.profile_url_raw for r in group}
        entry = {
            "usn": usn,
            "name": latest.full_name,
            "submissions": len(group),
            "kept_row": latest.row_number,
            "kept_timestamp": latest.timestamp.isoformat() if latest.timestamp else None,
            "dropped_rows": [r.row_number for r in group[:-1]],
            "profile_conflict": len(profile_ids) > 1,
        }
        report.duplicates_resolved.append(entry)
        if entry["profile_conflict"]:
            entry["profile_urls"] = [r.profile_url_raw for r in group]
            report.usn_conflicts.append(entry)
            log.warning(
                "USN %s submitted %d times with different profiles; keeping row %d (%s)",
                usn, len(group), latest.row_number, entry["kept_timestamp"],
            )
    return kept


def _reset_scrape_state(member: Member) -> None:
    member.last_attempt_at = None
    member.last_success_at = None
    member.last_error = None
    member.is_stale = False
    member.display_name = None
    member.member_since = None
    member.league = None
    member.platform_points = None
    member.avatar_url = None


def import_form(
    session: Session,
    form_path: str | Path,
    teams_path: str | Path | None = None,
) -> ImportReport:
    """Upsert members from a form export. Caller owns the transaction (commit/rollback)."""
    report = ImportReport()
    raw_rows, has_team_column = _read_form_rows(form_path, report)
    rows = _dedupe(raw_rows, report)
    report.unique_students = len(rows)

    team_map: dict[str, str] | None = None
    if teams_path:
        team_map = read_team_mapping(teams_path)
        report.team_source = "file"
    elif has_team_column:
        team_map = {}
        report.team_source = "column"
    if team_map is not None and has_team_column:
        # File entries win over the sheet column.
        for row in rows:
            if row.team and row.usn not in team_map:
                team_map[row.usn] = row.team

    teams_by_name = {t.name.lower(): t for t in session.scalars(select(Team))}

    def team_for(name: str) -> Team:
        team = teams_by_name.get(name.lower())
        if team is None:
            team = Team(name=name)
            session.add(team)
            teams_by_name[name.lower()] = team
        return team

    existing = {m.usn: m for m in session.scalars(select(Member))}
    seen: set[str] = set()

    for row in rows:
        seen.add(row.usn)
        link = check_profile_url(row.profile_url_raw)
        flags: list[str] = []
        ref = {"usn": row.usn, "name": row.full_name}

        if row.registration.strip().lower().startswith("no"):
            flags.append(FLAG_NOT_REGISTERED)
            report.not_registered.append(ref)
        if row.email and "@" not in row.email:
            flags.append(FLAG_INVALID_EMAIL)
            report.invalid_emails.append({**ref, "field": "email"})
        if row.college_email and "@" not in row.college_email:
            flags.append(FLAG_INVALID_COLLEGE_EMAIL)
            report.invalid_emails.append({**ref, "field": "college_email"})
        if any(d["usn"] == row.usn for d in report.duplicates_resolved):
            flags.append(FLAG_DUPLICATE_SUBMISSION)

        if link.status == MISSING_URL:
            report.missing_links.append(ref)
        elif link.status == INVALID_URL:
            report.invalid_links.append({**ref, "url": row.profile_url_raw})
        elif link.status == UNSUPPORTED_PLATFORM:
            report.unsupported_links.append({**ref, "url": row.profile_url_raw})

        member = existing.get(row.usn)
        if member is None:
            member = Member(usn=row.usn, link_status=link.status)
            session.add(member)
            existing[row.usn] = member
            report.created += 1
        else:
            report.updated += 1
            if not member.active:
                report.reactivated += 1
            url_changed = (member.profile_url or None) != (link.url or None)
            if url_changed:
                report.profile_changed += 1
                if member.id is not None:
                    session.query(Badge).filter(Badge.member_id == member.id).delete(
                        synchronize_session=False
                    )
                    session.expire(member, ["badges"])
                _reset_scrape_state(member)
                member.link_status = link.status
            elif member.link_status in IMPORT_STATUSES:
                member.link_status = link.status
            # else: keep the collector's last status (ok / not_found / ...)

        member.full_name = row.full_name
        member.email = row.email or None
        member.college_email = row.college_email or None
        member.year = row.year
        member.department = row.department or None
        member.registration_answer = row.registration or None
        member.submitted_at = row.timestamp
        member.profile_url_raw = row.profile_url_raw or None
        member.profile_url = link.url
        member.profile_id = link.profile_id
        member.platform = link.platform
        member.flags = flags
        member.active = True

        if team_map is not None:
            team_name = team_map.get(row.usn)
            member.team = team_for(team_name) if team_name else None

    for usn, member in existing.items():
        if usn not in seen and member.active:
            member.active = False
            report.deactivated += 1

    session.flush()

    # Same profile used by different students.
    by_profile: dict[str, list[Member]] = defaultdict(list)
    for member in existing.values():
        if member.active and member.profile_id:
            by_profile[member.profile_id].append(member)
    for profile_id, members in by_profile.items():
        if len(members) > 1:
            for member in members:
                member.flags = [*member.flags, FLAG_SHARED_PROFILE]
            report.shared_profiles.append(
                {"profile_id": profile_id, "usns": sorted(m.usn for m in members)}
            )
            log.warning("Profile %s is used by %s", profile_id, ", ".join(m.usn for m in members))

    active = [m for m in existing.values() if m.active]
    report.teams = len({m.team_id for m in active if m.team_id is not None})
    report.unassigned = [
        {"usn": m.usn, "name": m.full_name} for m in active if m.team_id is None
    ] if report.teams else []
    session.flush()
    return report


def format_report(report: ImportReport) -> str:
    lines = [
        "Import report",
        "=============",
        f"Rows read:              {report.total_rows}",
        f"Rows skipped:           {len(report.skipped_rows)}",
        f"Unique students:        {report.unique_students}",
        f"Created / updated:      {report.created} / {report.updated}",
        f"Reactivated:            {report.reactivated}",
        f"Deactivated (missing):  {report.deactivated}",
        f"Profile URL changed:    {report.profile_changed}",
        f"Duplicates resolved:    {len(report.duplicates_resolved)}",
        f"USN profile conflicts:  {len(report.usn_conflicts)}",
        f"Missing links:          {len(report.missing_links)}",
        f"Invalid links:          {len(report.invalid_links)}",
        f"Unsupported platform:   {len(report.unsupported_links)}",
        f"Shared profiles:        {len(report.shared_profiles)}",
        f"Not registered (No):    {len(report.not_registered)}",
        f"Invalid emails:         {len(report.invalid_emails)}",
        f"Teams ({report.team_source}):".ljust(24) + f"{report.teams}",
    ]
    if report.teams:
        lines.append(f"Students without team:  {len(report.unassigned)}")

    def section(title: str, items: list[dict[str, Any]], fmt) -> None:
        if items:
            lines.append("")
            lines.append(f"{title}:")
            lines.extend(f"  - {fmt(item)}" for item in items)

    section("Skipped rows", report.skipped_rows, lambda i: f"row {i['row']}: {i['reason']}")
    section(
        "USN conflicts (kept latest)",
        report.usn_conflicts,
        lambda i: f"{i['usn']} {i['name']}: kept row {i['kept_row']} ({i['kept_timestamp']}), "
        f"dropped rows {i['dropped_rows']}; urls={i.get('profile_urls')}",
    )
    section("Invalid links", report.invalid_links, lambda i: f"{i['usn']} {i['name']}: {i['url']}")
    section("Unsupported platform", report.unsupported_links, lambda i: f"{i['usn']} {i['name']}: {i['url']}")
    section("Missing links", report.missing_links, lambda i: f"{i['usn']} {i['name']}")
    section("Shared profiles", report.shared_profiles, lambda i: f"{i['profile_id']}: {', '.join(i['usns'])}")
    section("Not registered", report.not_registered, lambda i: f"{i['usn']} {i['name']}")
    section("Invalid emails", report.invalid_emails, lambda i: f"{i['usn']} {i['name']} ({i['field']})")
    section("Students without a team", report.unassigned, lambda i: f"{i['usn']} {i['name']}")
    return "\n".join(lines)
